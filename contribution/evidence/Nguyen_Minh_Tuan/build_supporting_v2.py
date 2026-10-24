"""Tạo data/Supporting_data_v2.xlsx — dữ liệu đã sửa + công thức tính mean/std + bảng nguồn (#1-#13)."""
import json, datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

S = json.load(open("/home/claude/stats.json"))
SRC = "/home/claude/d/Supporting data for fin assumptions.xlsx"
F = "Arial"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "666666"
H = PatternFill("solid", fgColor="D9E1F2"); Y = PatternFill("solid", fgColor="FFFF00")
wb = Workbook(); ws0 = wb.active; ws0.title = "README"

def cell(ws, ref, v, color=BLACK, bold=False, fmt=None, fill=None, italic=False, wrap=False):
    c = ws[ref]; c.value = v; c.font = Font(name=F, size=10, color=color, bold=bold, italic=italic)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c

def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = cell(ws, f"{get_column_letter(i)}{row}", n, bold=True, fill=H, wrap=True)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, text, note):
    cell(ws, "A1", text, bold=True); ws["A1"].font = Font(name=F, size=12, bold=True)
    cell(ws, "A2", note, color=GREY, italic=True)

# ---------------- VNIndex_fixed ----------------
ws = wb.create_sheet("VNIndex_fixed")
title(ws, "VN-Index đóng cửa tháng 05/2021–08/2026 (đã sửa lỗi ngày)", "Lỗi gốc: 'Mmm-YY' bị Sheets hiểu là Mmm-DD năm 2026 (VD 'May-21' → 21/05/2026) → năm thật nằm ở phần 'ngày'. Block hồi quy G6:O23 của file gốc hỏng do locale (R²=387, số dạng text) → 14,2% cũ không dùng được.")
header(ws, 4, ["stored_date (giá trị gốc, sai)", "Tháng thực", "VN-Index đóng cửa", "Lợi suất tháng", "Log return"], [22, 11, 16, 14, 12, 3, 30, 14, 16, 12, 60])
rows = sorted(S["vnindex_rows"], key=lambda r: (r["year"], r["month"]))
r0 = 5
for i, r in enumerate(rows):
    rr = r0 + i
    cell(ws, f"A{rr}", r["stored_date"], color=GREY)
    cell(ws, f"B{rr}", f"{r['year']}-{r['month']:02d}")
    cell(ws, f"C{rr}", r["close"], color=BLUE, fmt="#,##0.00")
    if i:
        cell(ws, f"D{rr}", f"=C{rr}/C{rr-1}-1", fmt="0.00%"); cell(ws, f"E{rr}", f"=LN(C{rr}/C{rr-1})", fmt="0.0000")
last = r0 + len(rows) - 1
summ = [("Số lợi suất tháng", f"=COUNT(D{r0+1}:D{last})", "0"), ("TB lợi suất tháng", f"=AVERAGE(D{r0+1}:D{last})", "0.00%"),
        ("Std tháng (mẫu)", f"=STDEV(D{r0+1}:D{last})", "0.00%"), ("TB năm = TB tháng × 12", "=H6*12", "0.00%"),
        ("Std năm = std tháng × √12", "=H7*SQRT(12)", "0.00%"), ("CAGR", f"=(C{last}/C{r0})^(12/H5)-1", "0.00%"),
        ("TB log return năm", f"=AVERAGE(E{r0+1}:E{last})*12", "0.00%")]
cell(ws, "G4", "Thống kê (công thức)", bold=True, fill=H); cell(ws, "H4", "Giá trị", bold=True, fill=H)
for i, (k, f, fmt) in enumerate(summ):
    cell(ws, f"G{5+i}", k); cell(ws, f"H{5+i}", f, fmt=fmt)
cell(ws, "G13", "Đối chiếu cuối năm", bold=True, fill=H); cell(ws, "H13", "Trong sheet", bold=True, fill=H)
cell(ws, "I13", "Báo chí", bold=True, fill=H); cell(ws, "J13", "Chênh", bold=True, fill=H); cell(ws, "K13", "Nguồn", bold=True, fill=H)
chk = [("2024-12", 1266.78, "Thanh Niên 31/12/2025: https://thanhnien.vn/vn-index-ket-thuc-nam-2025-voi-muc-tang-khung-gan-41-185251231154434459.htm"),
       ("2025-12", 1784.49, "Thanh Niên (như trên); VnEconomy: https://vneconomy.vn/vn-index-chot-2025-o-dinh-cao-lich-su-tang-ky-luc-9-nam.htm")]
for i, (m, v, src) in enumerate(chk):
    rr = 14 + i
    cell(ws, f"G{rr}", m); cell(ws, f"H{rr}", f'=INDEX($C${r0}:$C${last},MATCH("{m}",$B${r0}:$B${last},0))', fmt="#,##0.00")
    cell(ws, f"I{rr}", v, color=BLUE, fmt="#,##0.00"); cell(ws, f"J{rr}", f"=H{rr}-I{rr}", fmt="#,##0.00"); cell(ws, f"K{rr}", src, wrap=False)
cell(ws, "G16", "Tăng 2025 (sheet)"); cell(ws, "H16", "=H15/H14-1", fmt="0.00%"); cell(ws, "I16", 0.4087, color=BLUE, fmt="0.00%")
cell(ws, "J16", "=H16-I16", fmt="0.00%"); cell(ws, "K16", "VnEconomy: tăng 40,87% so với cuối 2024")
cell(ws, "G18", "Lưu ý: cửa sổ ~5 năm, chỉ số giá (chưa gồm cổ tức); đề xuất lấy mẫu per_year trong Monte Carlo.", color=GREY, italic=True)

# ---------------- Gold_SJC ----------------
ws = wb.create_sheet("Gold_SJC")
title(ws, "Giá vàng SJC (bán ra, cuối tháng) 01/2016–09/2026", "Nguồn: sheet 'Giá vàng SJC' của team (FiinPro). Ô xanh = số gốc; ô đen = công thức.")
header(ws, 4, ["Ngày", "Giá", "Lợi suất tháng"], [13, 14, 14, 3, 34, 14])
src = load_workbook(SRC, data_only=True)
gold = []
import unicodedata
gold_sheet = next(n for n in src.sheetnames if "sjc" in unicodedata.normalize("NFC", n).lower() or "vàng" in unicodedata.normalize("NFC", n).lower())
print("gold sheet:", repr(gold_sheet))
for row in src[gold_sheet].iter_rows(values_only=True):
    d = next((v for v in row if isinstance(v, (dt.datetime, dt.date))), None)
    if d is None:
        continue
    nums = [v for v in row if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 1000]
    if nums and dt.datetime(2016, 1, 1) <= (d if isinstance(d, dt.datetime) else dt.datetime(d.year, d.month, d.day)) <= dt.datetime(2026, 12, 31):
        gold.append((d, nums[0]))
gold.sort(key=lambda x: x[0])
assert len(gold) == S["gold"]["n_points"], (len(gold), S["gold"]["n_points"])
for i, (d, p) in enumerate(gold):
    rr = 5 + i
    cell(ws, f"A{rr}", d, fmt="yyyy-mm-dd"); cell(ws, f"B{rr}", p, color=BLUE, fmt="#,##0")
    if i:
        cell(ws, f"C{rr}", f"=B{rr}/B{rr-1}-1", fmt="0.00%")
gl = 4 + len(gold)
gs = [("Số lợi suất tháng", f"=COUNT(C6:C{gl})", "0"), ("TB năm = TB tháng × 12", f"=AVERAGE(C6:C{gl})*12", "0.00%"),
      ("Std năm = std tháng × √12", f"=STDEV(C6:C{gl})*SQRT(12)", "0.00%"), ("CAGR", f"=(B{gl}/B5)^(12/F5)-1", "0.00%")]
for i, (k, f, fmt) in enumerate(gs):
    cell(ws, f"E{5+i}", k); cell(ws, f"F{5+i}", f, fmt=fmt)
cell(ws, "E10", "Tăng giá Dec-to-Dec", bold=True, fill=H); cell(ws, "F10", "Lợi suất", bold=True, fill=H)
decs = {d.year: 5 + i for i, (d, _) in enumerate(gold) if d.month == 12}
k = 11
for y in sorted(decs):
    if y - 1 in decs:
        cell(ws, f"E{k}", str(y)); cell(ws, f"F{k}", f"=B{decs[y]}/B{decs[y-1]}-1", fmt="0.00%"); k += 1
cell(ws, f"E{k}", "Std năm (Dec-to-Dec)"); cell(ws, f"F{k}", f"=STDEV(F11:F{k-1})", fmt="0.00%")
cell(ws, f"E{k+1}", "Lưu ý: 2025 tăng bất thường; mean 14,9% là đặc thù giai đoạn — nên chạy độ nhạy.", color=GREY, italic=True)

# ---------------- Bond / Savings / CPI / Income ----------------
def series_sheet(name, ttl, note, cols, data, fmts, extra=None):
    ws = wb.create_sheet(name); title(ws, ttl, note)
    header(ws, 4, cols, [10] + [16] * (len(cols) - 1) + [3, 34, 14])
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            cell(ws, f"{get_column_letter(j+1)}{5+i}", v, color=(BLACK if j == 0 else BLUE), fmt=fmts[j])
    last = 4 + len(data)
    return ws, last

yrs = [str(int(y)) for y in S["bond"]["years"]]
ws, last = series_sheet("Bond_10Y", "Lợi suất TPCP Việt Nam 5 năm & 10 năm (cuối năm; 2026 = 25/09)", "Nguồn: sheet 'Return Bond' của team (FiinPro). Giả định nắm giữ đến đáo hạn.",
                        ["Năm", "TPCP 5 năm", "TPCP 10 năm"], list(zip(yrs, S["bond"]["y5"], S["bond"]["y10"])), ["@", "0.00%", "0.00%"])
for i, (k_, f) in enumerate([("TB 10 năm", f"=AVERAGE(C5:C{last})"), ("Std 10 năm", f"=STDEV(C5:C{last})"), ("TB 5 năm", f"=AVERAGE(B5:B{last})"), ("Std 5 năm", f"=STDEV(B5:B{last})")]):
    cell(ws, f"E{5+i}", k_); cell(ws, f"F{5+i}", f, fmt="0.00%")

ws, last = series_sheet("Savings_hist", "Lãi suất huy động VND cao nhất thị trường (cuối năm; 2026 = 25/09)", "Nguồn: sheet 'Lãi huy động' của team. Dùng mức CAO NHẤT → lệch lên so với lãi suất Big4.",
                        ["Năm", "Kỳ hạn 6–12 tháng", "Kỳ hạn >12 tháng"], list(zip(yrs, S["savings"]["s6_12"], S["savings"]["s_gt12"])), ["@", "0.00%", "0.00%"])
for i, (k_, f) in enumerate([("TB 6–12 tháng", f"=AVERAGE(B5:B{last})"), ("Std 6–12 tháng", f"=STDEV(B5:B{last})"), ("TB >12 tháng", f"=AVERAGE(C5:C{last})"), ("Std >12 tháng", f"=STDEV(C5:C{last})")]):
    cell(ws, f"E{5+i}", k_); cell(ws, f"F{5+i}", f, fmt="0.00%")

cpi_years = [str(y) for y in range(2016, 2026)]
ws, last = series_sheet("CPI_WB", "Lạm phát CPI Việt Nam 2016–2025 (World Bank)", "Nguồn: sheet 'Inflation WB' của team. Đề xuất cho living_cost_annual_growth_pct.",
                        ["Năm", "CPI (%/năm)"], [(y, v / 100) for y, v in zip(cpi_years, S["cpi"]["values"])], ["@", "0.00%"])
cell(ws, "E5", "TB 2016–2025"); cell(ws, "F5", f"=AVERAGE(B5:B{last})", fmt="0.00%")
cell(ws, "E6", "Std"); cell(ws, "F6", f"=STDEV(B5:B{last})", fmt="0.00%")

inc = S["income"]["levels"]
ws, last = series_sheet("Income_growth", "Thu nhập bình quân đầu người/tháng (VND) — KSMS 2019–2024", "Nguồn: NSO, Khảo sát mức sống dân cư 2024: https://www.nso.gov.vn/wp-content/uploads/2026/03/Sach-Khao-sat-muc-song-dan-cu-2024.pdf",
                        ["Năm", "Thu nhập (VND)", "Tăng trưởng"], [(str(int(y)), v) for y, v in inc], ["@", "#,##0"])
for rr in range(6, last + 1):
    cell(ws, f"C{rr}", f"=B{rr}/B{rr-1}-1", fmt="0.00%")
cell(ws, "E5", "TB tăng trưởng (đề xuất mean)"); cell(ws, "F5", f"=AVERAGE(C6:C{last})", fmt="0.00%")
cell(ws, "E6", "Std tăng trưởng (đề xuất std)"); cell(ws, "F6", f"=STDEV(C6:C{last})", fmt="0.00%")
cell(ws, "E8", "Cách team đang tính (SAI): std của MỨC thu nhập"); cell(ws, "F8", f"=STDEV(B5:B{last})", fmt="#,##0")
cell(ws, "E9", "→ đơn vị là VND, không phải % tăng trưởng; không dùng được cho Monte Carlo. n=5, có 2 năm COVID.", color=GREY, italic=True)

# ---------------- Property_CAGR (#3) ----------------
ws = wb.create_sheet("Property_CAGR")
title(ws, "Giá căn hộ sơ cấp Hà Nội (USD/m²) 2016–2025 và kiểm tra nhất quán giá/thuê (#3)", "Nguồn: sheet 'Lạm phát giá mua bds' của team (CafeF + Cushman & Wakefield). Ô vàng = tham số để thử.")
header(ws, 4, ["Năm", "USD/m²", "Tăng YoY"], [10, 12, 12, 3, 44, 14, 14, 14])
pp = S["property"]["prices"]
pyrs = sorted(pp, key=float)
for i, y in enumerate(pyrs):
    rr = 5 + i
    cell(ws, f"A{rr}", str(int(float(y)))); cell(ws, f"B{rr}", pp[y], color=BLUE, fmt="#,##0")
    if i:
        cell(ws, f"C{rr}", f"=B{rr}/B{rr-1}-1", fmt="0.0%")
pl = 4 + len(pyrs)
rows_ = [("CAGR 2016–2025", f"=(B{pl}/B5)^(1/9)-1"), ("CAGR 2016–2020", "=(B9/B5)^(1/4)-1"), ("CAGR 2020–2025", f"=(B{pl}/B9)^(1/5)-1"),
         ("TB tăng YoY", f"=AVERAGE(C6:C{pl})"), ("Std tăng YoY", f"=STDEV(C6:C{pl})")]
for i, (k_, f) in enumerate(rows_):
    cell(ws, f"E{5+i}", k_); cell(ws, f"F{5+i}", f, fmt="0.00%")
cell(ws, "E11", "Mốc neo dài hạn = CPI TB + tăng GDP thực/người", bold=True)
cell(ws, "E12", "CPI TB 2016–2025 (link CPI_WB)"); cell(ws, "F12", "=CPI_WB!F5", color=GREEN, fmt="0.00%")
cell(ws, "E13", "Tăng GDP thực/người (ước tính — Finance kiểm tra lại từ sheet 'GDP pc')"); cell(ws, "F13", 0.053, color=BLUE, fill=Y, fmt="0.00%")
cell(ws, "E14", "Mốc neo"); cell(ws, "F14", "=F12+F13", fmt="0.00%")
cell(ws, "E16", "Kiểm tra nhất quán sau N năm", bold=True, fill=H)
for j, h in enumerate(["19% (PA1 plan)", "CAGR 2016–25", "Mốc neo"]):
    cell(ws, f"{get_column_letter(6+j)}16", h, bold=True, fill=H)
cell(ws, "E17", "Tăng giá nhà g"); cell(ws, "F17", 0.19, color=BLUE, fill=Y, fmt="0.00%"); cell(ws, "G17", "=F5", fmt="0.00%"); cell(ws, "H17", "=F14", fmt="0.00%")
cell(ws, "E18", "Tăng giá thuê r (plan: 7%)"); cell(ws, "F18", 0.07, color=BLUE, fill=Y, fmt="0.00%"); cell(ws, "G18", "=F18", fmt="0.00%"); cell(ws, "H18", "=F18", fmt="0.00%")
cell(ws, "E19", "Số năm N"); cell(ws, "F19", 30, color=BLUE, fill=Y); cell(ws, "G19", "=F19"); cell(ws, "H19", "=F19")
cell(ws, "E20", "Lợi suất cho thuê ban đầu y0 (Vietnambiz: ~3% Q1/2025)"); cell(ws, "F20", 0.03, color=BLUE, fill=Y, fmt="0.00%"); cell(ws, "G20", "=F20", fmt="0.00%"); cell(ws, "H20", "=F20", fmt="0.00%")
for col in "FGH":
    cell(ws, f"{col}21", f"=(1+{col}17)^{col}19", fmt="#,##0.0\"x\""); cell(ws, f"{col}22", f"=(1+{col}17)^{col}19/(1+{col}18)^{col}19", fmt="#,##0.00\"x\"")
    cell(ws, f"{col}23", f"={col}20*(1+{col}18)^{col}19/(1+{col}17)^{col}19", fmt="0.00%")
cell(ws, "E21", "Giá nhà nhân lên (lần)"); cell(ws, "E22", "Tỷ lệ giá/thuê nhân lên (lần)"); cell(ws, "E23", "Lợi suất cho thuê ngụ ý sau N năm")
cell(ws, "E25", "Đọc kết quả: g = 19% và r = 7% trong 30 năm → lợi suất cho thuê ~0,1%/năm (không thực tế). Xem docs/KET_QUA_TODO.md #3.", color=GREY, italic=True)

# ---------------- Bảng nguồn ----------------
def source_sheet(name, ttl, rows, note=""):
    ws = wb.create_sheet(name); title(ws, ttl, note or "Mỗi dòng = 1 khẳng định + nguồn. Loại nguồn: Sơ cấp (trang chính thức) / Báo chí / Tổng hợp.")
    header(ws, 4, ["Hạng mục", "Nội dung / giá trị", "Thời điểm", "Loại nguồn", "Link", "Ghi chú độ tin cậy"], [24, 60, 14, 14, 70, 46])
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell(ws, f"{get_column_letter(j+1)}{5+i}", v, wrap=(j in (1, 5)))
    return ws

source_sheet("Nguon_vay_MB_VCB", "#1 thời hạn vay MB · #1b tuổi tối đa · #1c phí trả trước hạn · #2 LTV Vietcombank", [
    ("#1 MB Dream Home", "Vay tới 80%, thời hạn tới 35 năm (mô tả trang sản phẩm MB)", "tra cứu 09/2026", "Sơ cấp", "https://www.mbbank.com.vn/26/37/113/Chi-tiet/mua-nha-dat-tho-cu-chung-cu-du-an-cho-nguoi-tre-2024-6-5-17-58-30", "Gói riêng cho người trẻ — KHÁC gói CN Vạn Phúc (ân hạn 24 tháng)"),
    ("#1 MB Dream Home", "Dành cho khách hàng 23–40 tuổi", "28/03/2025", "Báo chí (Cổng TTĐT Chính phủ)", "https://baochinhphu.vn/mb-tien-phong-ho-tro-nguoi-tre-mua-nha-voi-goi-tin-dung-dream-home-102250328181355734.htm", ""),
    ("#1 MB Dream Home", "Hạn mức tối đa 5 tỷ; không trả gốc tới 5 năm đầu", "08/2024", "Báo chí", "https://thanhnien.vn/gioi-tre-hien-thuc-hoa-giac-mo-an-cu-voi-goi-vay-uu-dai-tu-mb-18524080416314239.htm", "Hồ sơ test (vay 5,7 tỷ) vượt hạn mức 5 tỷ → không áp dụng được gói này"),
    ("#1 MB Dream Home", "Ra mắt 03/2024 với thời hạn 30 năm", "03/2024", "Báo chí (VNBA)", "https://vnba.org.vn/vi/mb-ra-mat-goi-vay-mua-nha-uu-dai-dac-biet-danh-rieng-cho-khach-tre-13583.htm", "Thời hạn đã nâng lên 35 năm sau đó"),
    ("#1 MB gói khác", "Có gói MB cho vay tới 420 tháng", "không rõ", "Tổng hợp", "https://wiki.batdongsan.com.vn/wiki/cac-goi-vay-ngan-hang-mb-bank-809784", "Nguồn tổng hợp — cần MB xác nhận"),
    ("#1b tuổi tối đa MB", "Tuổi khi kết thúc khoản vay ≤ 70 (MB)", "không rõ", "Tổng hợp", "https://thebank.vn/vay-the-chap/cho-vay-mua-can-ho-nha-dat-du-an-ngan-hang-mbbank-40.html", "Mâu thuẫn với mốc 75 trong plan → xin MB Vạn Phúc xác nhận"),
    ("#1b đối chiếu", "PVcomBank: từ đủ 18 và không quá 75 tuổi khi kết thúc thời hạn cấp tín dụng", "tra cứu 09/2026", "Sơ cấp", "https://www.pvcombank.com.vn/tin-tuc/cam-nang-tai-chinh/vay-mua-nha-tra-gop", "Mốc 75 có thật nhưng ở ngân hàng khác"),
    ("#1b đối chiếu", "Techcombank: không quá 65 tuổi khi tất toán", "tra cứu 09/2026", "Sơ cấp", "https://techcombank.com/thong-tin/blog/dieu-kien-vay-the-chap", ""),
    ("#1c phí trả trước hạn", "Một gói MB: năm 1–2: 2,5%; năm 3: 2%; năm 4: 1,5%; năm 5: 1%; từ năm 6 miễn phí", "không rõ", "Tổng hợp", "https://wiki.batdongsan.com.vn/wiki/lai-suat-ngan-hang-mbbank-evr-803484", "Khác mức phẳng 1% × 5 năm trong to-do → đối chiếu văn bản Vạn Phúc"),
    ("#2 VCB Nhà Mới Thành Đạt", "Khách 20–35 tuổi; vay tối đa 70% giá trị nhà (50% nếu thuê mua); tới 40 năm; miễn trả gốc 5 năm đầu; tối đa 5 tỷ; thả nổi = LS huy động 24 tháng + 3%", "20/05/2025", "Báo chí (CafeF)", "https://cafef.vn/lai-suat-cho-vay-mua-nha-tai-vietcombank-vietinbank-agribank-bidv-thap-nhat-bao-nhieu-188250520101155192.chn", "Trang sản phẩm: https://www.vietcombank.com.vn/vi-VN/KHCN/SPDV/Vay/Nha-moi-thanh-dat"),
    ("#2 VCB vay mua nhà dự án", "Tối đa 70% nếu thế chấp chính căn nhà định mua; tới 100% nếu thế chấp tài sản khác; tối đa 20 năm", "cũ", "Tổng hợp", "https://banchungcusaigon.com.vn/vay-mua-bat-dong-san-tu-vietcombank.html", "Thông tin cũ — LTV phụ thuộc loại tài sản bảo đảm"),
    ("#2 VCB gói 20.000 tỷ", "Thời hạn tới 35 năm; tỷ lệ cho vay có thể tới 100% giá trị BĐS tuỳ hồ sơ", "2025", "Báo chí", "https://phapluat.suckhoedoisong.vn/vietcombank-mo-goi-tin-dung-20000-ty-dong-thi-truong-ha-lai-suat-vay-mua-nha-toi-1nam-284865.html", "Kết luận: 'LTV 100%' không sai ngân hàng mà là khác sản phẩm/khác TSBĐ"),
    ("#2 PDF plan trích", "20260729_Thong-bao-lai-suat-2026-thang-6.pdf — theo tên là thông báo LÃI SUẤT", "07/2026", "Sơ cấp", "https://www.vietcombank.com.vn/-/media/Project/VCB-Sites/VCB/Tin-tuc/2026/Vietcombank/Vietcombank_202607/Vietcombank_20260721/20260729_Thong-bao-lai-suat-2026-thang-6.pdf", "Không mở được trong phiên này; loại văn bản này thường không có LTV → khả năng trích nhầm"),
    ("#2 VCB lãi suất cơ sở", "Thông báo lãi suất cơ sở cho vay trung dài hạn (03/2026)", "16/03/2026", "Sơ cấp", "https://www.vietcombank.com.vn/vi-VN/Trang-thong-tin-dien-tu/Articles/2026/03/16/Thong-bao-LSCS-cho-vay-trung-dai-han", ""),
])
source_sheet("Lai_suat_cho_vay", "#6 Mặt bằng lãi suất cho vay (để kiểm tra std 0,75 điểm % của lãi thả nổi)", [
    ("LS cho vay mới BQ", "6,96%/năm cuối 2024; giữ quanh 6,96% đến cuối 11/2025", "11/2025", "Báo chí trích NHNN", "https://mekongasean.vn/lai-suat-cho-vay-binh-quan-giu-on-dinh-quanh-muc-696nam-50364.html", "Khoản vay MỚI, mọi kỳ hạn — không phải lãi vay mua nhà thả nổi"),
    ("LS cho vay mới BQ", "~6,34%/năm giữa 04/2025 (giảm 0,6 điểm % so với cuối 2024)", "04/2025", "Báo chí", "https://www.dnse.com.vn/senses/tin-tuc/lai-suat-cho-vay-ang-re-ky-luc-35086402", ""),
    ("LS cho vay BQ theo quý", "Thấp nhất 2020–2022 là 7,9%/năm (quý I/2022) — số liệu FiinPro", "04/2025", "Báo chí", "https://www.dnse.com.vn/senses/tin-tuc/lai-suat-cho-vay-ang-re-ky-luc-35086402", ""),
    ("LS cho vay BQ (cũ + mới)", "8,1–10,5%/năm (NHTM Nhà nước & cổ phần)", "07/2026", "Báo chí trích NHNN", "https://cafef.vn/so-lieu-lai-suat-cho-vay-lai-suat-huy-dong-moi-nhat-cua-he-thong-ngan-hang-188260721204319885.chn", "Bao gồm dư nợ cũ → biên rộng do khác sản phẩm, không phải biến động theo thời gian"),
    ("LS cho vay BQ (cũ + mới)", "8,4–10,7%/năm, tăng 0,2 điểm % so với tháng trước", "chưa xác minh ngày", "Báo chí", "https://vietnamfinance.vn/lai-suat-cho-vay-trung-binh-len-107-nam-thuc-te-khach-vay-cham-dinh-18-d150849.html", "Cần kiểm tra ngày đăng trước khi dùng"),
    ("Kiểm tra std theo thời gian", "Std LS huy động >12 tháng (mức cao nhất) 2016–2026 = 0,63 điểm % (sheet Savings_hist)", "", "Tính toán", "", "Lãi thả nổi = LS tham chiếu (thường là LS huy động) + biên → std ~0,6–0,8 ủng hộ 0,75 của plan. Cách (max−min)/4 trên khoảng cắt ngang KHÔNG đúng khái niệm"),
])
source_sheet("Gia_thue", "#7 Bằng chứng về tăng giá thuê căn hộ Hà Nội", [
    ("Giá thuê vs giá nhà", "Giá thuê chung cư ~ổn định quanh 13 triệu/căn trong 3 năm; lợi suất cho thuê giảm từ >4% (Q1/2023) xuống ~3% (Q1/2025); năm 2024 tốc độ tăng giá thuê chỉ ~một nửa tăng giá nhà", "04/2025", "Báo chí (Vietnambiz)", "https://vietnambiz.vn/loi-suat-cho-thue-chung-cu-dang-giam-dan-202541323742923.htm", "Bằng chứng gần nhất; không ủng hộ 7%/năm nếu giá thuê gần như đi ngang"),
    ("Giá thuê Q1/2025", "Giá thuê chung cư Hà Nội tăng nhẹ trong quý I/2025", "Q1/2025", "Tổng hợp (Batdongsan)", "https://wiki.batdongsan.com.vn/tin-tuc/gia-thue-chung-cu-Ha-Noi-tang-nhe-trong-quy-1-2025-cd-hn-832002", ""),
    ("Hiệp hội", "VARS: bài về giá thuê căn hộ tăng cao", "không rõ", "Hiệp hội", "https://vars.com.vn/tin-tuc/gia-thue-can-ho-tang-cao-n1897", "Chưa trích được con số cụ thể"),
    ("Kết luận cho mô hình", "Không có chuỗi giá thuê căn hộ thường của Savills/CBRE (chỉ có căn hộ dịch vụ). Nếu giữ 7% thì bắt buộc bật income growth; nếu rent growth > income growth kéo dài 30 năm, RENT tự động mất an toàn cuối kỳ", "", "Phân tích", "", "Xem docs/KET_QUA_TODO.md #7"),
])
source_sheet("Gia_BDS_thi_truong", "#3 Bằng chứng thị trường về tăng giá căn hộ Hà Nội", [
    ("Savills Q1/2025", "Giá sơ cấp +32% YoY; giá thứ cấp giảm nhẹ lần đầu sau 5 năm", "Q1/2025", "Báo cáo chuyên ngành", "https://vn.savills.com.vn/research_articles/164027/223536-0", "Tăng nóng tập trung ở sơ cấp"),
    ("Batdongsan qua Vietnambiz", "Giá chung cư Hà Nội +14% năm 2023, +39% năm 2024; 2025 tiếp tục tăng nhưng chậm lại", "10/2025", "Báo chí", "https://vietnambiz.vn/nhung-khu-vuc-tang-gia-chung-cu-manh-tai-ha-noi-202510985919392.htm", "Giai đoạn 2023–2025 là chu kỳ nóng, không đại diện 30 năm"),
    ("Dữ liệu của team", "CAGR 2016–2025 = 12,66%; 2016–2020 = 3,64%; 2020–2025 = 20,43% (sheet Property_CAGR)", "", "Tính toán", "", "18–20% chỉ khớp 5 năm gần nhất"),
])
ws = source_sheet("Tiet_kiem_snapshot", "#9 Lãi suất tiết kiệm hiện tại theo ngân hàng — MẪU ĐỂ TEAM ĐIỀN từ trang chính thức", [
    ("Báo chí 02/2026", "Vietcombank: kỳ hạn 12–18 tháng 5,2%; 24 tháng 5,3%", "22/02/2026", "Báo chí (CafeF)", "https://cafef.vn/nam-moi-gui-tiet-kiem-tai-agribank-vietcombank-bidv-vietinbank-cao-nhat-bao-nhieu-188260222210249665.chn", ""),
    ("Báo chí 03/2026", "Big4 tăng lãi suất: kỳ hạn 12–18 tháng từ 5,2% lên 5,9%", "03/2026", "Báo chí (Znews)", "https://znews.vn/vietcombank-vietinbank-bidv-agribank-cung-tang-lai-suat-tiet-kiem-post1638131.html", ""),
    ("Báo chí (ngày chưa rõ)", "Agribank/BIDV/VCB/VietinBank cùng niêm yết 6,8% kỳ hạn 12 tháng", "chưa rõ", "Báo chí", "https://giadinh.suckhoedoisong.vn/lai-suat-agribank-vietinbank-vietcombank-bidv-moi-nhat-gui-tiet-kiem-400-trieu-dong-vao-agribank-nhan-bao-nhieu-tien", "MÂU THUẪN với 2 dòng trên → không dùng khi chưa rõ ngày"),
], note="Các nguồn báo chí mâu thuẫn nhau → team điền bảng dưới từ biểu lãi suất CHÍNH THỨC của từng ngân hàng (ghi ngày tra cứu + link).")
cell(ws, "A10", "Bảng cần điền (ô vàng). Dòng 12 là VÍ DỤ định dạng — xoá khi điền.", bold=True)
for j, h in enumerate(["Ngân hàng", "6 tháng", "12 tháng", "18 tháng", "24 tháng", "Ngày tra cứu", "Link biểu lãi suất chính thức"]):
    cell(ws, f"{get_column_letter(j+1)}11", h, bold=True, fill=H)
ex = ["VÍ DỤ: Vietcombank", None, 0.052, 0.052, 0.053, "22/02/2026", "(ví dụ lấy từ CafeF — thay bằng link vietcombank.com.vn)"]
for j, v in enumerate(ex):
    cell(ws, f"{get_column_letter(j+1)}12", v, color=GREY, italic=True, fmt="0.00%" if 1 <= j <= 4 else None)
for i, b in enumerate(["Vietcombank", "BIDV", "VietinBank", "Agribank", "MB Bank"]):
    cell(ws, f"A{13+i}", b)
    for j in range(2, 8):
        cell(ws, f"{get_column_letter(j)}{13+i}", None, color=BLUE, fill=Y, fmt="0.00%" if j <= 5 else None)

# ---------------- Kenh_dau_tu (tổng hợp) ----------------
ws = wb.create_sheet("Kenh_dau_tu")
title(ws, "Tổng hợp 4 kênh đầu tư + độ lệch chuẩn Monte Carlo (giá trị đề xuất, chờ Coordinator + Finance duyệt)", "Ô xanh lá = link tới sheet tính toán. Sampling: per_path = rút 1 lần/đường; per_year = rút mỗi năm.")
header(ws, 4, ["Kênh / biến", "Mean (%/năm)", "Std (%/năm)", "Sampling", "Cửa sổ dữ liệu", "n", "Trạng thái", "Lưu ý"], [30, 13, 13, 11, 26, 6, 12, 60])
rows_ = [("Tiết kiệm (6–12 tháng, mức cao nhất)", "=Savings_hist!F5", "=Savings_hist!F6", "per_path", "cuối năm 2016–2026", 11, "Đề xuất", "Lệch lên (mức cao nhất); chờ snapshot #9"),
         ("Chứng khoán (VN-Index)", "=VNIndex_fixed!H8", "=VNIndex_fixed!H9", "per_year", "tháng 05/2021–08/2026", 63, "Đề xuất", "Cửa sổ ngắn; chưa gồm cổ tức"),
         ("Trái phiếu (TPCP 10 năm)", "=Bond_10Y!F5", "=Bond_10Y!F6", "per_path", "cuối năm 2016–2026", 11, "Đề xuất", "Giữ đến đáo hạn"),
         ("Vàng (SJC)", "=Gold_SJC!F6", "=Gold_SJC!F7", "per_year", "tháng 01/2016–09/2026", 128, "Đề xuất", "Giai đoạn tăng bất thường"),
         ("Tăng thu nhập (NSO)", "=Income_growth!F5", "=Income_growth!F6", "per_year", "2019–2024", 5, "Đề xuất", "Biến MỚI — bắt buộc nếu bật rent growth"),
         ("Tăng chi phí sinh hoạt (CPI)", "=CPI_WB!F5", 0, "—", "2016–2025", 10, "Đề xuất", "Biến MỚI; không random hoá"),
         ("Lãi thả nổi tham chiếu", None, 0.0075, "per_path", "", None, "Đã chốt", "Plan B8"),
         ("Tăng giá BĐS", 0.19, 0.04, "per_path", "", None, "Mean chờ #3", "Plan B8; xem Property_CAGR"),
         ("Tăng giá thuê", 0.07, 0.0125, "per_path", "", None, "Mean chờ #7", "Plan B8; xem Gia_thue")]
for i, r in enumerate(rows_):
    for j, v in enumerate(r):
        is_link = isinstance(v, str) and v.startswith("=")
        cell(ws, f"{get_column_letter(j+1)}{5+i}", v, color=(GREEN if is_link else (BLUE if j in (1, 2) and v is not None else BLACK)),
             fmt="0.00%" if j in (1, 2) else None, wrap=(j == 7))

# ---------------- LatLong_QA (#13) + Bathrooms ----------------
ws = wb.create_sheet("LatLong_QA")
title(ws, "#13 QA toạ độ PROJECT_MASTER", "3 dự án Hà Đông trùng hệt toạ độ (20.96069, 105.742059); toạ độ mới từ Google Maps/Places (tra cứu 29/09/2026) — cần team xác minh tòa cụ thể.")
header(ws, 4, ["project_id", "Tên dự án", "Lat cũ", "Long cũ", "Lat mới", "Long mới", "Lệch (km)", "Nguồn toạ độ mới"], [11, 26, 11, 12, 11, 12, 10, 60])
fx = json.load(open("/home/claude/finflow_repo/data/project_master_fixes.json", encoding="utf-8"))
pm = load_workbook("/home/claude/d/DATA.xlsx", read_only=True, data_only=True)["PROJECT_MASTER"]
it = pm.iter_rows(values_only=True); hdr = list(next(it)); idx = {h: i for i, h in enumerate(hdr) if h}
pmrows = {r[idx["project_id"]]: r for r in it if r and r[idx["project_id"]]}
for i, pid in enumerate(["HD072", "HD073", "HD074"]):
    rr = 5 + i; r = pmrows[pid]
    vals = [pid, r[idx["project_name"]], r[idx["latitude"]], r[idx["longitude"]], fx[pid]["latitude"], fx[pid]["longitude"]]
    for j, v in enumerate(vals):
        cell(ws, f"{get_column_letter(j+1)}{rr}", v, color=(BLUE if j >= 2 else BLACK), fmt="0.000000" if j >= 2 else None)
    cell(ws, f"G{rr}", f"=2*6371*ASIN(SQRT(SIN(RADIANS(E{rr}-C{rr})/2)^2+COS(RADIANS(C{rr}))*COS(RADIANS(E{rr}))*SIN(RADIANS(F{rr}-D{rr})/2)^2))", fmt="0.00")
    cell(ws, f"H{rr}", fx[pid]["source"])
bw = [(pid, r[idx["project_name"]], r[idx["beltway_zone"]]) for pid, r in pmrows.items() if isinstance(r[idx["beltway_zone"]], (int, float))]
cell(ws, "A10", f"beltway_zone lưu dạng SỐ (cần chuẩn hoá thành chuỗi '4'): {len(bw)} dự án", bold=True)
for i, (pid, nm, v) in enumerate(bw):
    cell(ws, f"A{11+i}", pid); cell(ws, f"B{11+i}", nm); cell(ws, f"C{11+i}", v, color=BLUE)
k = 12 + len(bw)
cell(ws, f"A{k}", "Công thức bathrooms (Google Sheets, cột J = bathrooms; CHƯA chạy thử trên Sheet thật):", bold=True)
cell(ws, f"A{k+1}", '=ARRAYFORMULA(IF(A2:A="","",IF((TRIM(J2:J)="")+(UPPER(TRIM(J2:J))="N/A"),"Không rõ",J2:J)))')
ws[f"A{k+1}"].data_type = "s"
cell(ws, f"A{k+2}", "OR() không chạy theo từng phần tử trong ARRAYFORMULA → dùng phép cộng (+). Plan có 2 công thức cho ra 2 nhãn khác nhau → cần chốt 1 nhãn. Hiện có 9 ô 'N/A' dạng text + 4 ô trống.", color=GREY, italic=True)

# ---------------- README ----------------
lines = ["Supporting_data_v2.xlsx — dữ liệu hỗ trợ FIN_ASSUMPTIONS đã sửa (FinFlow, 29/09/2026)", "",
         "Quy ước: chữ xanh dương = số gốc nhập từ nguồn; chữ đen = công thức; chữ xanh lá = link sang sheet khác; nền vàng = ô cần team điền/thử.",
         "Mọi mean/std đều là CÔNG THỨC trên dữ liệu gốc → sửa dữ liệu thì kết quả tự cập nhật.", "",
         "VNIndex_fixed — #10: sửa lỗi ngày, tính lại mean/std, đối chiếu cuối năm với báo chí",
         "Gold_SJC — #12: lợi suất vàng; Bond_10Y — #11; Savings_hist — #9 (chuỗi lịch sử)",
         "Income_growth — #8: std TĂNG TRƯỞNG thu nhập (team đang lấy std của MỨC thu nhập)", "CPI_WB — đề xuất tăng chi phí sinh hoạt",
         "Property_CAGR — #3: CAGR theo cửa sổ + kiểm tra nhất quán giá/thuê 30 năm", "Nguon_vay_MB_VCB — #1, #1b, #1c, #2",
         "Lai_suat_cho_vay — #6; Gia_thue — #7; Gia_BDS_thi_truong — #3; Tiet_kiem_snapshot — #9 (mẫu điền)",
         "Kenh_dau_tu — bảng tổng hợp mean/std/sampling cho FIN_ASSUMPTIONS và Monte Carlo", "LatLong_QA — #13 + công thức bathrooms"]
for i, t in enumerate(lines, 1):
    cell(ws0, f"A{i}", t, bold=(i == 1))
ws0.column_dimensions["A"].width = 120
wb.save("/home/claude/finflow_repo/data/Supporting_data_v2.xlsx")
print("saved; gold points", len(gold), "; beltway numeric", len(bw))
