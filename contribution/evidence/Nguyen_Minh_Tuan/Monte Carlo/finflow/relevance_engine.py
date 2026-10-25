"""B2 — MUA_CALC / THUE_CALC relevance engine (FINAL PLAN §B2, bản cập nhật 2026-09-25 thang 100, Spec §3).

Luồng: bộ lọc cứng (bedrooms=2, size_band khớp đúng — size_band là THAM SỐ của từng lần gọi) → TARGET = listing của
project_id user chọn ở Step 2 → chấm điểm giá/quận/vành đai (45/30/25) → xếp tier THEO SỐ TIÊU CHÍ KHỚP
(3→HIGH, 2→MEDIUM, 1→LOW, 0→WARNING; tránh lỗi trộn thang điểm cũ/mới) → trong tier xếp theo Haversine tăng dần →
ngưỡng hiển thị TARGET + HIGH/MEDIUM + top 5 LOW + top 3 WARNING. THUE_CALC dùng chung 100% logic (giá = thuê/m²).

Diễn giải đã chọn (cần team xác nhận — docs/KET_QUA_TODO.md, câu hỏi Q1-Q3):
  - Bộ lọc Step 1 chỉ thu hẹp DANH SÁCH TÒA để chọn ở Step 2 (filter_buildings); tập candidate của B2 = mọi listing
    cùng size band ở cả 3 quận (nếu lọc theo quận thì tiêu chí 'cùng quận' luôn đúng → vô nghĩa).
  - Mốc giá TARGET = trung vị giá/m² các listing của tòa TARGET trong size band; nếu tòa không có listing → trung vị
    của toàn bộ candidate cùng size band.
  - N (quy tắc N<5) = số listing trong pool sau bộ lọc cứng (gồm cả listing TARGET).
"""
import math
import statistics

EARTH_RADIUS_KM = 6371.0
DEFAULT_WEIGHTS = {"price": 45, "district": 30, "beltway": 25}  # CHOT (to-do B2: 45/30/25, không phải giả lập)
PRICE_TOLERANCE = {"Small": 0.05, "Medium": 0.08, "Large": 0.10}
N_MIN_FOR_PRICE = 5
DISPLAY_CAPS = {"TARGET": None, "HIGH": None, "MEDIUM": None, "LOW": 5, "WARNING": 3}
LEVEL_ORDER = ("TARGET", "HIGH", "MEDIUM", "LOW", "WARNING")
LEVEL_BY_MATCHES = {3: "HIGH", 2: "MEDIUM", 1: "LOW", 0: "WARNING"}
WARNING_BANNER = "Các căn mức WARNING chỉ mang tính tham khảo lỏng — không khớp tiêu chí nào ngoài size band."
EPS = 1e-9


def normalize_beltway(v):
    """Chuẩn hoá vành đai: 4.0 (số) / '4.0' → '4'; '2.5-3' giữ nguyên; None/NaN → None."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and math.isnan(v):
            return None
        return str(int(v)) if float(v).is_integer() else f"{v:g}"
    s = str(v).strip()
    try:
        f = float(s)
    except ValueError:
        return s
    return str(int(f)) if f.is_integer() else s


def haversine_km(lat1, lon1, lat2, lon2):
    """distance_km = 2 × 6371 × asin(√(sin²(Δlat/2) + cos(lat1)·cos(lat2)·sin²(Δlon/2))) — đúng công thức plan."""
    if None in (lat1, lon1, lat2, lon2):
        return None
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi, dlmb = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    h = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(math.sqrt(min(1.0, h)))


def score_table(weights=None):
    """Bảng điểm suy ra từ trọng số (khớp bảng 'Thang mới 100' của plan cập nhật và bảng 80 của plan gốc)."""
    w = DEFAULT_WEIGHTS if weights is None else weights
    p, d, b = w["price"], w["district"], w["beltway"]
    pairs, singles = (d + b, b + p, d + p), (b, d, p)
    if max(singles) >= min(pairs) or max(pairs) >= p + d + b:
        raise ValueError("Trọng số làm chồng lấn tier (1 tiêu chí ≥ 2 tiêu chí) — không hợp lệ")
    return {"HIGH": p + d + b, "district+beltway": d + b, "beltway+price": b + p, "district+price": d + p,
            "beltway_only": b, "district_only": d, "price_only": p,
            "MEDIUM_range": (min(pairs), p + d + b - 1), "LOW_range": (min(singles), min(pairs) - 1),
            "cap_when_n_lt_min": d + b}


def listing_price_m2(listing, mode):
    """BUY: price_m2_vnd (hoặc tổng/diện tích). RENT: rent_monthly_vnd / area_m2 (RENT_RAW không có cột giá/m²)."""
    area = listing.get("area_m2")
    if mode == "buy":
        v = listing.get("price_m2_vnd")
        if v is None and listing.get("total_price_vnd") and area:
            v = listing["total_price_vnd"] / area
    else:
        v = listing["rent_monthly_vnd"] / area if listing.get("rent_monthly_vnd") and area else None
    return float(v) if v is not None and v > 0 else None


def _fmt_ref(value, mode):
    return f"{value / 1e6:.1f} triệu/m²" if mode == "buy" else f"{value / 1e3:.0f} nghìn/m²/tháng"


def run_relevance_engine(listings, projects, target_project_id, size_band, mode, *, weights=None, tolerances=None,
                         n_min=N_MIN_FOR_PRICE, bedrooms=2, display_caps=None):
    """Chạy 1 lần MUA_CALC (mode='buy') hoặc THUE_CALC (mode='rent') cho 1 size band."""
    if mode not in ("buy", "rent"):
        raise ValueError("mode phải là 'buy' hoặc 'rent'")
    tol_map = PRICE_TOLERANCE if tolerances is None else tolerances
    if size_band not in tol_map:
        raise ValueError(f"size_band phải là một trong {list(tol_map)}")
    if target_project_id not in projects:
        raise ValueError(f"TARGET project_id '{target_project_id}' không có trong PROJECT_MASTER")
    w = DEFAULT_WEIGHTS if weights is None else weights
    caps = DISPLAY_CAPS if display_caps is None else display_caps
    table = score_table(w)
    tgt = projects[target_project_id]
    tgt_d, tgt_b = tgt.get("district_old"), normalize_beltway(tgt.get("beltway_zone"))

    pool, excluded_no_price = [], 0
    for L in listings:
        if L.get("size_band") != size_band:
            continue
        try:
            if float(L.get("bedrooms")) != float(bedrooms):
                continue
        except (TypeError, ValueError):
            continue
        pm2 = listing_price_m2(L, mode)
        if pm2 is None:
            excluded_no_price += 1
            continue
        pool.append((L, pm2))
    n_pool = len(pool)
    price_enabled = n_pool >= n_min
    target_rows = [(L, p) for L, p in pool if L["project_id"] == target_project_id]
    cand_rows = [(L, p) for L, p in pool if L["project_id"] != target_project_id]
    if target_rows:
        ref, ref_source = statistics.median(p for _, p in target_rows), "TARGET"
    elif cand_rows:
        ref, ref_source = statistics.median(p for _, p in cand_rows), "MEDIAN_SIZE_BAND"
    else:
        ref, ref_source = None, None
    tol = tol_map[size_band]
    ref_label = "mốc TARGET" if ref_source == "TARGET" else "mốc trung vị size band"

    rows = []
    for L, p in target_rows:
        rows.append({**_base(L, p, mode), "relevance_level": "TARGET", "relevance_score": None, "n_matched": None,
                     "matched": None, "price_diff_pct": (p / ref - 1) if ref else None, "distance_km": 0.0,
                     "warning_reason": f"TARGET — đúng tòa bạn chọn ({tgt.get('project_name', target_project_id)})."})
    for L, p in cand_rows:
        proj = projects.get(L["project_id"], {})
        c_d = L.get("district_old") or proj.get("district_old")
        c_b = normalize_beltway(L.get("beltway_zone") if L.get("beltway_zone") is not None else proj.get("beltway_zone"))
        diff = (p / ref - 1) if ref else None
        m_price = bool(price_enabled and diff is not None and abs(diff) <= tol + EPS)
        m_dist = c_d is not None and c_d == tgt_d
        m_belt = c_b is not None and c_b == tgt_b
        n_match = int(m_price) + int(m_dist) + int(m_belt)
        score = w["price"] * m_price + w["district"] * m_dist + w["beltway"] * m_belt
        level = LEVEL_BY_MATCHES[n_match]
        ok, ko = [], []
        if not price_enabled:
            ko.append(f"không đủ dữ liệu để so sánh giá (N={n_pool} < {n_min})")
        elif m_price:
            ok.append(f"giá/m² lệch {diff * 100:+.1f}% so với {ref_label} {_fmt_ref(ref, mode)} (trong ±{tol * 100:g}%)")
        else:
            ko.append(f"giá/m² lệch {diff * 100:+.1f}% so với {ref_label} {_fmt_ref(ref, mode)} (ngoài ±{tol * 100:g}%)")
        (ok if m_dist else ko).append(f"cùng quận ({c_d})" if m_dist else f"khác quận ({c_d} ≠ {tgt_d})")
        (ok if m_belt else ko).append(f"cùng vành đai ({c_b})" if m_belt else f"khác vành đai ({c_b} ≠ {tgt_b})")
        reason = "Khớp: " + ("; ".join(ok) if ok else "không tiêu chí nào") + ". Không khớp: " + ("; ".join(ko) if ko else "—") + "."
        rows.append({**_base(L, p, mode), "relevance_level": level, "relevance_score": score, "n_matched": n_match,
                     "matched": {"price": m_price, "district": m_dist, "beltway": m_belt}, "price_diff_pct": diff,
                     "distance_km": haversine_km(tgt.get("latitude"), tgt.get("longitude"),
                                                 proj.get("latitude"), proj.get("longitude")),
                     "warning_reason": reason})
    # Kiểm tra điểm nằm đúng khoảng của tier (bảo vệ khỏi lỗi trộn thang điểm 80/100)
    for r in rows:
        lvl, sc = r["relevance_level"], r["relevance_score"]
        if lvl == "HIGH":
            assert sc == table["HIGH"], r
        elif lvl == "MEDIUM":
            assert table["MEDIUM_range"][0] <= sc <= table["MEDIUM_range"][1], r
        elif lvl == "LOW":
            assert table["LOW_range"][0] <= sc <= table["LOW_range"][1], r
        elif lvl == "WARNING":
            assert sc == 0, r
        if not price_enabled and sc is not None:
            assert sc <= table["cap_when_n_lt_min"], r
    order = {lvl: i for i, lvl in enumerate(LEVEL_ORDER)}
    rows.sort(key=lambda r: (order[r["relevance_level"]], r["distance_km"] is None,
                             r["distance_km"] if r["distance_km"] is not None else 0.0, r["listing_id"]))
    counter = {}
    for r in rows:
        counter[r["relevance_level"]] = counter.get(r["relevance_level"], 0) + 1
        r["candidate_rank"] = counter[r["relevance_level"]]
    display = [r for r in rows if caps.get(r["relevance_level"]) is None or r["candidate_rank"] <= caps[r["relevance_level"]]]
    return {
        "meta": {"mode": mode, "size_band": size_band, "target_project_id": target_project_id,
                 "target_district": tgt_d, "target_beltway": tgt_b, "reference_price_m2": ref,
                 "reference_source": ref_source, "tolerance": tol, "n_pool": n_pool, "n_min": n_min,
                 "price_scoring_enabled": price_enabled, "excluded_no_price": excluded_no_price,
                 "weights": dict(w), "counts": counter},
        "candidates": rows,
        "display": display,
        "show_warning_banner": any(r["relevance_level"] == "WARNING" for r in display),
        "warning_banner": WARNING_BANNER,
    }


def _base(L, p, mode):
    return {"listing_id": L["listing_id"], "project_id": L["project_id"], "project_name": L.get("project_name"),
            "district_old": L.get("district_old"), "ward_old": L.get("ward_old"),
            "beltway_zone": normalize_beltway(L.get("beltway_zone")), "area_m2": L.get("area_m2"),
            "size_band": L.get("size_band"), "price_m2": p,
            "price_total": L.get("total_price_vnd") if mode == "buy" else L.get("rent_monthly_vnd")}


def mua_calc(buy_listings, projects, target_project_id, size_band_buy, **kw):
    return run_relevance_engine(buy_listings, projects, target_project_id, size_band_buy, "buy", **kw)


def thue_calc(rent_listings, projects, target_project_id, size_band_rent, **kw):
    return run_relevance_engine(rent_listings, projects, target_project_id, size_band_rent, "rent", **kw)


def run_step2(data, target_project_id, size_band_buy, size_band_rent, **kw):
    """Step 2: 2 lời gọi ĐỘC LẬP (size band BUY, size band RENT), cùng TARGET project_id."""
    from .disclaimers import size_band_disclaimers
    buy = mua_calc(data["buy"], data["projects"], target_project_id, size_band_buy, **kw)
    rent = thue_calc(data["rent"], data["projects"], target_project_id, size_band_rent, **kw)
    return {"buy": buy, "rent": rent, "disclaimers": size_band_disclaimers(size_band_buy, size_band_rent)}


def filter_buildings(projects, district_old, ward_old=None, beltway_zone=None):
    """Step 1: thu hẹp danh sách tòa cho dropdown Step 2 (quận bắt buộc, phường/vành đai tuỳ chọn)."""
    if not district_old:
        raise ValueError("Step 1: quận cũ là trường bắt buộc")
    bz = normalize_beltway(beltway_zone) if beltway_zone else None
    out = [pid for pid, p in projects.items() if p.get("district_old") == district_old
           and (not ward_old or p.get("ward_old") == ward_old)
           and (bz is None or normalize_beltway(p.get("beltway_zone")) == bz)]
    return sorted(out, key=lambda pid: projects[pid].get("project_name") or pid)


def load_data_xlsx(path, coordinate_fixes=None):
    """Đọc DATA.xlsx (giá trị cache của Google Sheets) → {'projects', 'buy', 'rent'}; áp bản sửa toạ độ nếu có."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)

    def rows_of(name):
        it = wb[name].iter_rows(values_only=True)
        header = [h for h in next(it)]
        keep = [i for i, h in enumerate(header) if h is not None]
        out = []
        for r in it:
            if r is None or all(r[i] is None for i in keep):
                continue
            out.append({header[i]: r[i] for i in keep})
        return out

    projects = {}
    for r in rows_of("PROJECT_MASTER"):
        if not r.get("project_id"):
            continue
        projects[r["project_id"]] = {"project_id": r["project_id"], "project_name": r.get("project_name"),
                                     "district_old": r.get("district_old"), "ward_old": r.get("ward_old"),
                                     "ward_2025": r.get("ward_2025"), "beltway_zone": normalize_beltway(r.get("beltway_zone")),
                                     "latitude": r.get("latitude"), "longitude": r.get("longitude")}
    for pid, fix in (coordinate_fixes or {}).items():
        if pid in projects:
            projects[pid].update({"latitude": fix["latitude"], "longitude": fix["longitude"]})
    buy = [dict(r, beltway_zone=normalize_beltway(r.get("beltway_zone"))) for r in rows_of("BUY_RAW") if r.get("listing_id")]
    rent = [dict(r, beltway_zone=normalize_beltway(r.get("beltway_zone"))) for r in rows_of("RENT_RAW") if r.get("listing_id")]
    return {"projects": projects, "buy": buy, "rent": rent}


def qa_project_master(projects):
    """QA cho A1: toạ độ trùng hệt nhau giữa các dự án khác nhau, toạ độ thiếu, vành đai lạ."""
    by_ll = {}
    for pid, p in projects.items():
        if p.get("latitude") is None or p.get("longitude") is None:
            continue
        by_ll.setdefault((round(p["latitude"], 6), round(p["longitude"], 6)), []).append(pid)
    return {"duplicate_coordinates": [v for v in by_ll.values() if len(v) > 1],
            "missing_coordinates": [pid for pid, p in projects.items() if p.get("latitude") is None or p.get("longitude") is None],
            "beltway_values": sorted({p.get("beltway_zone") for p in projects.values() if p.get("beltway_zone")})}
